from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "tmp" / "doc-report" / "source.docx"
OUTPUT = ROOT / "output" / "documents" / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_final_v9.docx"
ASSETS = ROOT / "slides" / "customer-analytics" / "assets"
CHATBOT = ROOT / "tmp" / "doc-report" / "chatbot.png"
TECH_ICONS = ROOT / "tmp" / "doc-report" / "tech-icons"
FIGURES = [
    "Hình 1. Kiến trúc tổng thể của hệ thống",
    "Hình 2. Màn hình đăng nhập",
    "Hình 3. Dashboard sau xác thực",
    "Hình 4. Trang chủ phía khách hàng",
    "Hình 5. Giao diện xác thực",
    "Hình 6. Giao diện trợ lý 3CS AI",
    "Hình 7. Dashboard tổng quan",
    "Hình 8. Danh sách phân khúc",
    "Hình 9. Danh sách khách hàng ưu tiên",
]
TABLES = [
    "Bảng 1. Phân công và đánh giá thành viên",
    "Bảng 2. Công nghệ sử dụng",
    "Bảng 3. Vai trò người dùng",
    "Bảng 4. Các thành phần chính của hệ thống",
    "Bảng 5. Luồng hoạt động chính",
    "Bảng 6. Môi trường phát triển",
    "Bảng 7. Các chỉ số đánh giá mô hình",
    "Bảng 8. Agent Skills hỗ trợ phát triển",
    "Bảng 9. Kiểm thử chức năng",
]
table_number = 0


def set_font(run, name: str = "Times New Roman", size: float | None = None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is None:
        no_wrap = OxmlElement("w:noWrap")
        tc_pr.append(no_wrap)


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def set_col_widths(table, widths_cm):
    for column, width in zip(table.columns, widths_cm, strict=False):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm, strict=False):
            cell.width = Cm(width)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.0)

    for style_name, size, align in (
        ("Heading 1", 15, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13.5, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12.5, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)


def add_field(paragraph, instruction: str, *, size: float = 10):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=size)


def add_toc_field(paragraph):
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở tài liệu."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])
    set_font(run, size=12)


def add_hyperlink(paragraph, text: str, url: str, size: float = 10.5):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = paragraph.add_run(text)
    set_font(run, size=size)
    run.font.color.rgb = RGBColor(5, 99, 193)
    run.font.underline = True
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_body_section(doc: Document):
    cover = doc.sections[0]
    cover.page_width = Cm(21)
    cover.page_height = Cm(29.7)
    cover.top_margin = Cm(1.8)
    cover.bottom_margin = Cm(1.8)
    cover.left_margin = Cm(2.2)
    cover.right_margin = Cm(2.0)
    cover.header_distance = Cm(0.7)
    cover.footer_distance = Cm(0.7)

    for paragraph in cover.header.paragraphs:
        paragraph.text = ""
    for paragraph in cover.footer.paragraphs:
        paragraph.text = ""

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.85)

    section.header.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.text = "IE221 - Kỹ thuật lập trình Python"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_font(run, size=9)
        run.font.color.rgb = RGBColor(89, 89, 89)

    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")

    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), "1")


def add_text(doc, text: str, *, bold_lead: str | None = None, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.0) if indent else Cm(0)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, size=13, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_font(rest, size=13)
    else:
        run = p.add_run(text)
        set_font(run, size=13)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(f"•  {item}")
        set_font(run, size=12.5)


def add_chapter(doc, title: str):
    doc.add_page_break()
    p = doc.add_paragraph(title, style="Heading 1")
    p.paragraph_format.space_after = Pt(12)
    return p


def add_heading(doc, text: str, level: int = 2):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    figure_number = next((i for i, caption in enumerate(FIGURES, 1) if caption == text), None)
    if figure_number is not None:
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), str(100 + figure_number))
        bookmark_start.set(qn("w:name"), f"figure_{figure_number}")
        p._p.append(bookmark_start)
    run = p.add_run(text)
    set_font(run, size=10)
    run.italic = True
    if figure_number is not None:
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), str(100 + figure_number))
        p._p.append(bookmark_end)
    return p


def add_table_caption(doc, text: str, number: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(200 + number))
    start.set(qn("w:name"), f"table_{number}")
    p._p.append(start)
    run = p.add_run(text)
    set_font(run, size=10, bold=True)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(200 + number))
    p._p.append(end)


def add_single_figure(doc, path: Path, caption: str, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_two_figures(doc, left: Path, left_caption: str, right: Path, right_caption: str):
    add_single_figure(doc, left, left_caption, 13.5)
    add_single_figure(doc, right, right_caption, 13.5)


def add_table(doc, headers, rows, widths_cm=None, font_size=10.5, alignments=None):
    global table_number
    table_number += 1
    add_table_caption(doc, TABLES[table_number - 1], table_number)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths_cm:
        set_col_widths(table, widths_cm)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers, strict=False):
        set_cell_shading(cell, "D9E2F3")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.right_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_font(run, size=font_size, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, (cell, value) in enumerate(
            zip(row.cells, values, strict=False)
        ):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, "F7F9FC")
            p = cell.paragraphs[0]
            p.alignment = (
                alignments[column_index]
                if alignments and column_index < len(alignments)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.right_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_icon_label(paragraph, icon: str, label: str, *, icon_width=0.36, size=8.5):
    picture_run = paragraph.add_run()
    picture_run.add_picture(str(TECH_ICONS / f"{icon}.png"), width=Cm(icon_width))
    text_run = paragraph.add_run(f" {label}")
    set_font(text_run, size=size)


def add_icon_line(cell, icon: str, label: str, *, icon_width=0.36, size=8.5):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    add_icon_label(paragraph, icon, label, icon_width=icon_width, size=size)
    return paragraph


def add_architecture_table(doc):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_col_widths(table, [4.0, 0.8, 4.2, 0.8, 5.2])
    remove_table_borders(table)
    titles = ("Frontend", "↔", "Backend", "↔", "Dữ liệu và dịch vụ")
    for index, title in enumerate(titles):
        cell = table.rows[0].cells[index]
        set_cell_margins(cell, top=150, start=100, bottom=150, end=100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(3)
        if title == "↔":
            run = p.add_run(title)
            set_font(run, size=16)
            continue
        set_cell_shading(cell, "F7F9FC")
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "6")
            tag.set(qn("w:color"), "BFBFBF")
            borders.append(tag)
        tc_pr.append(borders)
        run = p.add_run(title)
        set_font(run, size=11, bold=True)

        if title == "Frontend":
            add_icon_line(cell, "react", "React")
            add_icon_line(cell, "typescript", "TypeScript")
            add_icon_line(cell, "vite", "Vite")
        elif title == "Backend":
            add_icon_line(cell, "fastapi", "FastAPI")
            add_icon_line(cell, "python", "Python")
        else:
            add_icon_line(cell, "postgresql", "PostgreSQL")
            add_icon_line(cell, "python", "ML")
            p3 = cell.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.paragraph_format.first_line_indent = Cm(0)
            p3.paragraph_format.space_after = Pt(0)
            p3.add_run().add_picture(str(TECH_ICONS / "cloudinary.png"), width=Cm(1.45))
            p3.add_run("   ")
            p3.add_run().add_picture(str(TECH_ICONS / "groq.png"), width=Cm(1.0))
    add_caption(doc, "Hình 1. Kiến trúc tổng thể của hệ thống")


def build():
    doc = Document(str(SOURCE))

    body = doc._element.body
    cover_table = doc.tables[0]._tbl
    final_sect_pr = body.sectPr
    for child in list(body):
        if child is cover_table or child is final_sect_pr:
            continue
        body.remove(child)
    if body.index(cover_table) != 0:
        body.remove(cover_table)
        body.insert(0, cover_table)

    configure_styles(doc)
    add_body_section(doc)

    # Body page 1: automatic contents from Heading 1-3.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(14)
    set_font(p.add_run("MỤC LỤC"), size=15, bold=True)
    toc_paragraph = doc.add_paragraph()
    add_toc_field(toc_paragraph)

    doc.add_page_break()
    doc.add_paragraph("DANH MỤC HÌNH ẢNH", style="Heading 1")
    for figure_number, caption in enumerate(FIGURES, 1):
        entry = doc.add_paragraph()
        entry.alignment = WD_ALIGN_PARAGRAPH.LEFT
        entry.paragraph_format.first_line_indent = Cm(0)
        entry.paragraph_format.left_indent = Cm(0)
        entry.paragraph_format.right_indent = Cm(0)
        entry.paragraph_format.line_spacing = 1.2
        entry.paragraph_format.space_after = Pt(4)
        entry.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        set_font(entry.add_run(f"{caption}\t"), size=12)
        add_field(entry, f"PAGEREF figure_{figure_number} \\h", size=12)

    doc.add_paragraph("DANH MỤC BẢNG", style="Heading 1")
    for number, caption in enumerate(TABLES, 1):
        entry = doc.add_paragraph()
        entry.alignment = WD_ALIGN_PARAGRAPH.LEFT
        entry.paragraph_format.first_line_indent = Cm(0)
        entry.paragraph_format.left_indent = Cm(0)
        entry.paragraph_format.right_indent = Cm(0)
        entry.paragraph_format.line_spacing = 1.2
        entry.paragraph_format.space_after = Pt(4)
        entry.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        set_font(entry.add_run(f"{caption}\t"), size=12)
        add_field(entry, f"PAGEREF table_{number} \\h", size=12)

    # Body page 2: contributions and thanks.
    doc.add_page_break()
    doc.add_paragraph("BẢNG PHÂN CÔNG VÀ ĐÁNH GIÁ THÀNH VIÊN", style="Heading 1")
    assignment_table = add_table(
        doc,
        ["Họ và tên", "MSSV", "Nội dung phụ trách", "Đánh giá"],
        [
            ["Võ Thị Hương Giang", "25410196", "Phân tích nghiệp vụ và thiết kế Backend API", "Hoàn thành"],
            ["Phan Huy Hoàng", "25410219", "Phát triển Frontend và tích hợp API", "Hoàn thành"],
            ["Nguyễn Minh Nam", "25410259", "Phát triển Backend API và xử lý dữ liệu", "Hoàn thành"],
        ],
        [4.8, 2.7, 4.7, 3.2],
        9,
        [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    for cell_index in (0, 1, 3):
        set_cell_no_wrap(assignment_table.rows[0].cells[cell_index])
    for row in assignment_table.rows[1:]:
        for cell_index in (0, 1, 3):
            set_cell_no_wrap(row.cells[cell_index])
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "LỜI CẢM ƠN", 2)
    add_text(
        doc,
        "Nhóm 6 xin chân thành cảm ơn ThS. Nghi Hoàng Khoa đã hướng dẫn môn Kỹ thuật lập trình Python, góp ý về cách tổ chức chương trình và định hướng nhóm hoàn thiện đề tài. Những kiến thức về Python, lập trình hướng đối tượng và tổ chức mã nguồn là nền tảng để nhóm xây dựng hệ thống theo hướng có cấu trúc, dễ kiểm tra và tiếp tục mở rộng.",
    )
    add_text(
        doc,
        "Nhóm cũng cảm ơn Trường Đại học Công nghệ Thông tin đã tạo môi trường học tập và cung cấp các điều kiện cần thiết cho quá trình thực hiện đồ án. Do thời gian và dữ liệu thực nghiệm còn giới hạn, báo cáo khó tránh khỏi thiếu sót; nhóm mong nhận được góp ý để tiếp tục cải thiện hệ thống.",
    )

    # Body page 3: foreword.
    doc.add_page_break()
    doc.add_paragraph("LỜI NÓI ĐẦU", style="Heading 1")
    add_text(
        doc,
        "Hoạt động bán hàng trên nền tảng số tạo ra dữ liệu từ hồ sơ khách hàng, đơn hàng, sản phẩm và các lượt tương tác. Khi dữ liệu nằm rời rạc, người quản trị khó xác định khách hàng cần ưu tiên, khó theo dõi xu hướng mua sắm và khó đánh giá khả năng mua lại. Đề tài “Hệ thống phân tích dữ liệu khách hàng và dự đoán tiềm năng mua hàng” được thực hiện để giải quyết bài toán đó trong phạm vi một đồ án môn học.",
    )
    add_text(
        doc,
        "Sản phẩm gồm website khách hàng, trang quản trị và dịch vụ API. Backend Python tiếp nhận dữ liệu nghiệp vụ, tính điểm tiềm năng theo hành vi mua hàng, huấn luyện mô hình dự đoán mua lại và cung cấp kết quả cho giao diện. Frontend hỗ trợ duyệt sản phẩm, ghi nhận lượt quan tâm, quản lý dữ liệu và theo dõi các báo cáo phân tích.",
    )
    add_text(
        doc,
        "Báo cáo tập trung vào cách nhóm áp dụng Python trong kiến trúc backend, tổ chức lớp và module, xử lý dữ liệu, tính điểm, huấn luyện mô hình và kiểm thử. Kết quả dự đoán chỉ đóng vai trò hỗ trợ ra quyết định; độ tin cậy phụ thuộc vào chất lượng dữ liệu, cửa sổ quan sát và điều kiện đánh giá mô hình.",
    )

    # Body page 4: chapter 1.
    add_chapter(doc, "CHƯƠNG 1. GIỚI THIỆU")
    add_heading(doc, "1.1. Lý do chọn đề tài")
    add_text(doc, "Dữ liệu khách hàng chỉ tạo ra giá trị khi được tổng hợp và diễn giải thành thông tin có thể hành động. Một hệ thống kết hợp dữ liệu giao dịch với hành vi tương tác giúp doanh nghiệp nhận biết khách hàng có giá trị, theo dõi nhu cầu và xây dựng danh sách chăm sóc phù hợp. Đây cũng là bài toán phù hợp để vận dụng Python vào API, xử lý dữ liệu và mô hình dự đoán trong cùng một sản phẩm.")
    add_heading(doc, "1.2. Mục tiêu")
    add_bullets(
        doc,
        [
            "Quản lý tập trung thông tin khách hàng, sản phẩm, danh mục và đơn hàng.",
            "Tính điểm tiềm năng từ lịch sử mua sắm và mức độ tương tác.",
            "Huấn luyện, đánh giá và triển khai mô hình dự đoán khả năng mua lại.",
            "Trình bày kết quả qua dashboard, danh sách ưu tiên và báo cáo có thể xuất dữ liệu.",
            "Tổ chức mã nguồn Python theo lớp rõ ràng, dễ kiểm thử và bảo trì.",
        ],
    )
    add_heading(doc, "1.3. Phạm vi và giới hạn")
    add_text(doc, "Đề tài tập trung vào luồng bán hàng, dữ liệu khách hàng, hành vi xem sản phẩm và phân tích phục vụ quản trị. Hệ thống dùng dữ liệu trong PostgreSQL và hỗ trợ nhập dữ liệu theo hợp đồng định dạng xác định. Phần dự đoán xét khả năng mua lại trong một khoảng thời gian cấu hình. Một số màn hình hoặc dữ liệu minh họa vẫn phục vụ mục đích trình diễn; kết quả chưa thay thế quy trình đánh giá kinh doanh trên dữ liệu sản xuất quy mô lớn.")
    add_heading(doc, "1.4. Cấu trúc báo cáo")
    add_text(doc, "Báo cáo gồm sáu chương: giới thiệu; cơ sở lý thuyết và công nghệ; phân tích và thiết kế; hiện thực và cài đặt; kết quả và đánh giá; kết luận và hướng phát triển.")

    # Body pages 5-6: chapter 2.
    add_chapter(doc, "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ")
    add_heading(doc, "2.1. Phân tích khách hàng")
    add_text(doc, "Phân tích khách hàng kết hợp dữ liệu mô tả, giao dịch và tương tác để trả lời ba câu hỏi: khách hàng đã mua gần đây hay chưa, mua với tần suất nào và mang lại giá trị bao nhiêu. Đề tài dùng nhóm chỉ số RFM làm nền tảng, sau đó bổ sung mức độ tương tác để phản ánh tín hiệu quan tâm trước khi phát sinh đơn hàng.")
    add_heading(doc, "2.2. Điểm tiềm năng")
    add_text(doc, "Mỗi thành phần được chuẩn hóa về thang điểm tối đa 5. Trọng số mặc định gồm Recency 35%, Frequency 30%, Monetary 20% và Interaction 15%. Điểm tổng hợp được quy đổi về thang 100 theo công thức:")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("Potential Score = (R × 0,35 + F × 0,30 + M × 0,20 + I × 0,15) × 20")
    set_font(run, size=11.5, bold=True)
    add_text(doc, "Hệ thống phân loại từ 80 điểm trở lên là mức Cao, từ 60 đến dưới 80 là Tiềm năng và dưới 60 là Bình thường. Nếu không có đơn hàng hợp lệ trong cửa sổ phân tích, hệ thống giữ trạng thái thiếu dữ liệu thay vì thay giá trị thiếu bằng 0.")
    add_heading(doc, "2.3. Dự đoán khả năng mua lại")
    add_text(doc, "Mô hình sử dụng các đặc trưng như recency, frequency, monetary, giá trị đơn trung bình, chu kỳ mua, điểm tương tác, độ đa dạng sản phẩm và điểm đánh giá. Hệ thống hỗ trợ Logistic Regression và Random Forest được hiện thực bằng Python. Việc phê duyệt mô hình dựa trên PR-AUC, Lift@Top10 và Precision@Top10 để đánh giá đúng bài toán có tỷ lệ chuyển đổi không cân bằng.")
    add_heading(doc, "2.4. Nguyên tắc đánh giá")
    add_text(doc, "Mô hình được huấn luyện từ cửa sổ đặc trưng và dự đoán cho một khoảng thời gian tương lai. Phiên bản mới không tự động phục vụ người dùng ngay sau huấn luyện; hệ thống chỉ cho phép triển khai khi vượt qua các ngưỡng đánh giá đã cấu hình. Cách tiếp cận này tách quá trình thử nghiệm khỏi quá trình sử dụng kết quả.")

    doc.add_page_break()
    add_heading(doc, "2.5. Python và lập trình hướng đối tượng")
    add_text(doc, "Python 3.13 là ngôn ngữ chính của backend. FastAPI đảm nhiệm lớp HTTP và kiểm tra dữ liệu đầu vào thông qua Pydantic. SQLAlchemy cung cấp ánh xạ đối tượng quan hệ và giao tiếp bất đồng bộ với PostgreSQL. Mã nguồn chia theo domain, application, infrastructure và presentation; lớp nghiệp vụ phụ thuộc vào giao diện repository thay vì truy cập trực tiếp cơ sở dữ liệu.")
    add_heading(doc, "2.6. Công nghệ sử dụng")
    add_table(
        doc,
        ["Nhóm", "Công nghệ", "Vai trò"],
        [
            ["Frontend", "React, TypeScript, Vite", "Giao diện khách hàng và quản trị"],
            ["Backend", "Python, FastAPI, Pydantic", "API, xác thực, nghiệp vụ và phân tích"],
            ["Dữ liệu", "PostgreSQL, SQLAlchemy, Alembic", "Lưu trữ, truy vấn và quản lý phiên bản lược đồ"],
            ["Dịch vụ", "Cloudinary, Groq Cloud", "Lưu trữ ảnh sản phẩm và trợ lý AI"],
            ["Công cụ", "uv, pnpm, pytest, Vitest", "Quản lý phụ thuộc, chạy và kiểm thử"],
        ],
        [3.0, 5.2, 7.2],
        10.5,
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_heading(doc, "2.7. Quy trình xử lý")
    add_text(doc, "Luồng phân tích bắt đầu từ dữ liệu khách hàng, đơn hàng và tương tác. Backend kiểm tra dữ liệu, tính các đặc trưng, chấm điểm tiềm năng, huấn luyện hoặc nạp mô hình đã triển khai, sau đó trả kết quả cho dashboard và danh sách ưu tiên. Mỗi kết quả lưu kèm phiên bản cấu hình hoặc phiên bản mô hình để có thể kiểm tra lại.")

    # Body pages 7-9: chapter 3.
    add_chapter(doc, "CHƯƠNG 3. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
    add_heading(doc, "3.1. Đối tượng sử dụng")
    add_table(
        doc,
        ["Vai trò", "Nhu cầu chính", "Quyền tiêu biểu"],
        [
            ["Khách hàng", "Duyệt sản phẩm, xem chi tiết và tạo tín hiệu quan tâm", "Đọc sản phẩm, ghi nhận lượt xem"],
            ["Quản trị viên", "Quản lý dữ liệu, theo dõi dashboard và vận hành mô hình", "Quản lý tài khoản, dữ liệu, phân tích và mô hình"],
        ],
        [3.0, 7.0, 5.4],
        10.5,
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_heading(doc, "3.2. Yêu cầu chức năng")
    add_bullets(
        doc,
        [
            "Đăng ký, đăng nhập, làm mới phiên, đặt lại mật khẩu và phân quyền theo vai trò.",
            "Quản lý khách hàng, tài khoản, sản phẩm, danh mục và đơn hàng.",
            "Ghi nhận lượt xem sản phẩm và liên kết tài khoản với hồ sơ khách hàng.",
            "Tính điểm tiềm năng, phân khúc khách hàng và xuất dữ liệu phân tích.",
            "Huấn luyện, liệt kê, phê duyệt và triển khai phiên bản mô hình.",
            "Tạo danh sách khách hàng ưu tiên và hỗ trợ truy vấn qua trợ lý 3CS AI.",
        ],
    )
    add_heading(doc, "3.3. Yêu cầu phi chức năng")
    add_text(doc, "Hệ thống cần phản hồi rõ ràng khi dữ liệu không hợp lệ, bảo vệ phiên đăng nhập, không ghi log mật khẩu hoặc token, giới hạn quyền theo từng endpoint và tách cấu hình bí mật khỏi mã nguồn. Giao diện phải thích ứng với màn hình khác nhau, hỗ trợ sáng/tối và biểu diễn trạng thái đang tải bằng skeleton.")

    doc.add_page_break()
    add_heading(doc, "3.4. Kiến trúc tổng thể")
    add_architecture_table(doc)
    add_text(doc, "Frontend gọi các API có kiểu dữ liệu rõ ràng và quản lý trạng thái server bằng TanStack Query. Backend nhận request tại lớp presentation, gọi use case ở lớp application, thao tác với entity và repository interface ở domain, sau đó truy cập PostgreSQL qua implementation bất đồng bộ ở infrastructure. Cloudinary lưu ảnh sản phẩm; Groq hỗ trợ chức năng trò chuyện AI.")
    add_heading(doc, "3.5. Các thành phần chính")
    add_table(
        doc,
        ["Thành phần", "Trách nhiệm"],
        [
            ["Identity", "Xác thực, phiên làm việc, OTP, tài khoản và phân quyền"],
            ["Customer và Order", "Hồ sơ khách hàng, lịch sử mua và chỉ số giao dịch"],
            ["Product", "Sản phẩm, danh mục, ảnh và sự kiện xem sản phẩm"],
            ["Analytics", "RFM, điểm tiềm năng, dashboard, mô hình và danh sách ưu tiên"],
            ["Import Data", "Đọc workbook, kiểm tra hợp đồng dữ liệu và ghi dữ liệu theo giao dịch"],
        ],
        [4.2, 11.2],
        10.5,
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_page_break()
    add_heading(doc, "3.6. Luồng hoạt động chính")
    add_table(
        doc,
        ["Bước", "Xử lý"],
        [
            ["1", "Khách hàng truy cập danh mục hoặc xem chi tiết sản phẩm."],
            ["2", "Hệ thống ghi nhận tương tác và liên kết với hồ sơ khách hàng."],
            ["3", "Backend tổng hợp giao dịch, RFM và điểm tương tác trong cửa sổ phân tích."],
            ["4", "Điểm tiềm năng và xác suất mua lại được tính theo cấu hình hoặc mô hình đã triển khai."],
            ["5", "Dashboard và danh sách ưu tiên trình bày kết quả cho quản trị viên."],
        ],
        [2.4, 13.0],
        10.5,
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "3.7. Thiết kế đăng nhập và phân quyền")
    add_two_figures(doc, ASSETS / "login.png", "Hình 2. Màn hình đăng nhập", ASSETS / "dashboard.png", "Hình 3. Dashboard sau xác thực")
    add_text(doc, "Access token được giữ trong bộ nhớ ở frontend; refresh token được backend đặt trong cookie HTTP-only. Khi nhiều request cùng nhận lỗi 401, client chỉ thực hiện một request làm mới phiên dùng chung rồi thử lại các request đang chờ. Route guard và permission dependency cùng tham gia kiểm soát truy cập ở hai phía.")

    # Body pages 10-13: chapter 4.
    add_chapter(doc, "CHƯƠNG 4. HIỆN THỰC VÀ CÀI ĐẶT")
    add_heading(doc, "4.1. Môi trường phát triển")
    add_table(
        doc,
        ["Thành phần", "Yêu cầu"],
        [
            ["JavaScript", "Node.js 22 trở lên, pnpm 10 trở lên"],
            ["Python", "Python 3.13, quản lý môi trường bằng uv"],
            ["Cơ sở dữ liệu", "PostgreSQL có thể truy cập từ backend"],
            ["Chạy cục bộ", "Frontend cổng 4000, backend cổng 8000"],
        ],
        [4.0, 11.4],
        10.5,
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "4.2. Cấu trúc mã nguồn")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(8)
    structure = (
        "customer-analytics-platform/\n"
        "├── frontend/src/        # React, routes, components, hooks và API client\n"
        "├── backend/src/customer_analytics/\n"
        "│   ├── app/features/    # Các lát cắt nghiệp vụ\n"
        "│   ├── app/shared/      # Schema và xử lý lỗi dùng chung\n"
        "│   └── core/            # Database, middleware và unit of work\n"
        "├── backend/migrations/  # Alembic\n"
        "└── backend/tests/       # Pytest"
    )
    run = p.add_run(structure)
    set_font(run, name="Consolas", size=9.5)
    add_heading(doc, "4.3. Khởi chạy")
    add_text(doc, "Sau khi cấu hình backend/.env và kết nối PostgreSQL, lệnh pnpm setup cài phụ thuộc cho toàn monorepo. Lệnh pnpm dev chạy đồng thời frontend và backend. API cung cấp Swagger tại /docs để kiểm tra endpoint trong quá trình phát triển.")

    doc.add_page_break()
    add_heading(doc, "4.4. Hiện thực xác thực và quản lý tài khoản")
    add_text(doc, "Backend kiểm tra email, mật khẩu, trạng thái tài khoản và quyền trước khi thực thi use case. Access token có thời hạn ngắn; refresh session có thể bị thu hồi khi đăng xuất hoặc đặt lại mật khẩu. Quy trình quên mật khẩu sử dụng OTP sáu chữ số, giới hạn thời gian và chỉ cho phép dùng một lần.")
    add_heading(doc, "4.5. Hiện thực danh mục sản phẩm")
    add_text(doc, "API sản phẩm trả danh sách phân trang, chi tiết sản phẩm và tối đa bốn sản phẩm liên quan cùng danh mục. Ảnh tải lên được kiểm tra MIME type và giới hạn dung lượng trước khi chuyển sang Cloudinary. Lượt xem sản phẩm được lưu thành interaction_type product_view để phục vụ phân tích hành vi.")
    add_single_figure(doc, ASSETS / "client-home.png", "Hình 4. Trang chủ phía khách hàng", 9.5)
    add_single_figure(doc, ASSETS / "login.png", "Hình 5. Giao diện xác thực", 9.5)
    add_heading(doc, "4.6. Trợ lý AI")
    add_text(doc, "Trợ lý 3CS AI sử dụng Groq Cloud để phản hồi câu hỏi về dữ liệu và nghiệp vụ trong phạm vi được cho phép. Giao diện trò chuyện hỗ trợ màn hình nhỏ, chế độ toàn màn hình và trạng thái đóng/mở có chuyển động nhất quán.")
    add_single_figure(doc, CHATBOT, "Hình 6. Giao diện trợ lý 3CS AI", 8.0)

    heading = add_heading(doc, "4.7. Tính điểm tiềm năng")
    heading.paragraph_format.page_break_before = True
    add_text(doc, "Use case tính điểm đọc các dải điểm và trọng số từ cấu hình, tổng hợp hành vi trong khoảng thời gian lựa chọn, chuẩn hóa từng thành phần rồi lưu kết quả kèm phiên bản cấu hình. Khi thiếu thành phần Interaction, trọng số các thành phần còn lại được chuẩn hóa lại; khi thiếu dữ liệu giao dịch cốt lõi, kết quả được đánh dấu thiếu dữ liệu.")
    add_heading(doc, "4.8. Huấn luyện và quản lý mô hình")
    add_text(doc, "Request huấn luyện nhận loại mô hình, cửa sổ đặc trưng, khoảng dự đoán và ngày phân tích tùy chọn. Nếu ngày phân tích để trống, backend chọn mốc dữ liệu phù hợp. Sau huấn luyện, artifact và metric được đăng ký theo version. Quản trị viên chỉ triển khai phiên bản ở trạng thái APPROVED; endpoint dự đoán trả lỗi dịch vụ nếu chưa có mô hình được triển khai.")
    add_table(
        doc,
        ["Metric", "Ý nghĩa trong hệ thống"],
        [
            ["PR-AUC", "Đánh giá chất lượng xếp hạng trong dữ liệu chuyển đổi không cân bằng"],
            ["Lift@Top10", "So sánh tỷ lệ chuyển đổi của nhóm ưu tiên cao nhất với toàn bộ tập"],
            ["Precision@Top10", "Tỷ lệ khách hàng mua lại trong nhóm 10% được xếp hạng cao nhất"],
        ],
        [5.0, 10.4],
        10.5,
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "4.9. Danh sách ưu tiên")
    add_text(doc, "Danh sách ưu tiên kết hợp điểm tiềm năng và xác suất mua lại. Khách hàng có điểm từ 80 trở lên hoặc xác suất vượt ngưỡng cấu hình được đưa vào danh sách; thứ tự ưu tiên dựa trên cả hai đại lượng để hỗ trợ hoạt động chăm sóc.")

    doc.add_page_break()
    add_heading(doc, "4.10. Giao diện quản trị và phân tích")
    add_single_figure(doc, ASSETS / "dashboard.png", "Hình 7. Dashboard tổng quan", 15.5)
    add_text(doc, "Dashboard cho phép lọc theo thời gian, phân khúc, mức tiềm năng, nhóm sản phẩm và nhân viên phụ trách. Các chỉ số doanh thu, đơn hàng, khách hàng và giá trị đơn trung bình được trình bày cùng biểu đồ xu hướng, phân bố khách hàng và dữ liệu chất lượng.")
    add_two_figures(doc, ASSETS / "segments.png", "Hình 8. Danh sách phân khúc", ASSETS / "priority.png", "Hình 9. Danh sách khách hàng ưu tiên")

    doc.add_page_break()
    add_heading(doc, "4.11. Công cụ AI và Agent Skills hỗ trợ phát triển")
    add_text(
        doc,
        "Trong quá trình thực hiện, nhóm sử dụng các công cụ AI để hỗ trợ đọc mã nguồn, đề xuất cấu trúc, triển khai giao diện, rà soát lỗi, chạy kiểm thử và biên soạn tài liệu. Mọi thay đổi do AI đề xuất đều được nhóm kiểm tra, chỉnh sửa và xác nhận bằng các bước typecheck, lint hoặc test phù hợp trước khi tích hợp.",
    )
    add_heading(doc, "4.11.1. Nền tảng AI và môi trường agent", 3)

    ai_tools = [
        (
            "OMP",
            "Hỗ trợ làm việc với coding agent trong môi trường gắn với IDE, dùng để đọc mã nguồn, chỉnh sửa và kiểm tra thay đổi.",
            "https://omp.sh/",
        ),
        (
            "Orca",
            "Môi trường phát triển dành cho agent, hỗ trợ quản lý workspace, terminal, trình duyệt, diff và các tác vụ phát triển.",
            "https://www.onorca.dev/",
        ),
        (
            "Codex Pro 20x",
            "Trợ lý lập trình AI được sử dụng để hỗ trợ triển khai tính năng, gỡ lỗi, viết kiểm thử, rà soát giao diện và hoàn thiện tài liệu.",
            "https://openai.com/codex/",
        ),
    ]
    for name, description, url in ai_tools:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run(f"• {name}: "), size=12.5, bold=True)
        set_font(p.add_run(description + " "), size=12.5)
        add_hyperlink(p, url, url, size=11.5)

    add_heading(doc, "4.11.2. Agent Skills sử dụng", 3)
    add_text(
        doc,
        "Các skill được cài đặt từ hệ sinh thái skills.sh và dùng như hướng dẫn quy trình cho từng nhóm công việc. Danh sách skill chính gồm:",
    )
    add_table(
        doc,
        ["Skill", "Mục đích sử dụng"],
        [
            ["fastapi", "Áp dụng quy ước FastAPI, Pydantic, dependency và API bất đồng bộ."],
            ["fastapi-templates", "Tham khảo cấu trúc dự án, dependency injection và xử lý lỗi."],
            ["shadcn", "Xây dựng và chuẩn hóa component giao diện dựa trên shadcn/ui."],
            ["typescript-advanced-types", "Tăng an toàn kiểu cho schema, API client và component TypeScript."],
            ["vercel-composition-patterns", "Tổ chức component React theo hướng dễ tái sử dụng và mở rộng."],
            ["vercel-react-best-practices", "Rà soát hiệu năng, hooks và các thực hành React phù hợp."],
        ],
        [5.2, 10.2],
        10,
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    source_p = doc.add_paragraph()
    source_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    source_p.paragraph_format.first_line_indent = Cm(0)
    source_p.paragraph_format.space_before = Pt(5)
    set_font(source_p.add_run("Nguồn Agent Skills: "), size=11.5, bold=True)
    add_hyperlink(source_p, "https://skills.sh/", "https://skills.sh/", size=11.5)

    # Body pages 14-15: chapter 5.
    add_chapter(doc, "CHƯƠNG 5. KẾT QUẢ VÀ ĐÁNH GIÁ")
    add_heading(doc, "5.1. Kết quả đạt được")
    add_text(doc, "Đồ án hoàn thành một luồng xuyên suốt từ trải nghiệm khách hàng đến quản trị và phân tích. Người dùng có thể duyệt sản phẩm và đăng nhập; quản trị viên quản lý dữ liệu, theo dõi dashboard, xem phân khúc, vận hành mô hình và truy cập danh sách ưu tiên. Backend cung cấp cấu trúc lớp rõ ràng, xử lý lỗi tập trung và hợp đồng API có kiểu dữ liệu cụ thể.")
    add_heading(doc, "5.2. Kiểm thử chức năng")
    add_table(
        doc,
        ["Chức năng", "Tình huống kiểm tra", "Kết quả mong đợi", "Trạng thái"],
        [
            ["Đăng nhập", "Thông tin hợp lệ hoặc sai mật khẩu", "Tạo phiên hoặc trả lỗi xác thực rõ ràng", "Đạt"],
            ["Phân quyền", "Tài khoản không có quyền gọi API quản trị", "Từ chối truy cập", "Đạt"],
            ["Danh mục", "Tải danh sách và chi tiết sản phẩm", "Dữ liệu phân trang và sản phẩm liên quan", "Đạt"],
            ["Tương tác", "Người dùng xem chi tiết sản phẩm", "Ghi nhận sự kiện product_view", "Đạt"],
            ["Điểm tiềm năng", "Có và thiếu dữ liệu RFM", "Tính đúng mức hoặc trả trạng thái thiếu dữ liệu", "Đạt"],
            ["Huấn luyện", "Ngày phân tích để trống", "Payload không gửi analysis_date", "Đạt"],
            ["Video và slide", "Mở, đóng, tải lại tại một slide", "Phát/dừng đúng và giữ số slide trên URL", "Đạt"],
        ],
        [3.4, 4.0, 5.4, 2.6],
        9.2,
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    add_heading(doc, "5.3. Kiểm tra chất lượng mã nguồn")
    add_text(doc, "Frontend sử dụng TypeScript để kiểm tra kiểu, ESLint và Prettier để duy trì quy ước mã nguồn, Vitest cho kiểm thử hành vi. Backend dùng Ruff, mypy và pytest. Ở lần kiểm tra cuối, frontend đạt 113/113 bài kiểm thử trên 32 tệp và backend đạt 71/71 bài kiểm thử. Các bài kiểm thử tập trung vào use case, schema, ánh xạ lỗi và endpoint quan trọng.")

    doc.add_page_break()
    add_heading(doc, "5.4. Ưu điểm")
    add_bullets(
        doc,
        [
            "Kiến trúc backend phân lớp, tách domain khỏi hạ tầng và framework.",
            "Luồng xác thực và refresh token hạn chế việc lưu token dài hạn ở trình duyệt.",
            "Điểm tiềm năng có thể truy vết nhờ lưu thành phần, trọng số và phiên bản cấu hình.",
            "Vòng đời mô hình tách huấn luyện, phê duyệt và triển khai.",
            "Giao diện hỗ trợ responsive, dark mode, trạng thái tải và các màn hình phân tích chính.",
        ],
    )
    add_heading(doc, "5.5. Hạn chế")
    add_bullets(
        doc,
        [
            "Chất lượng dự đoán phụ thuộc vào lượng đơn hàng và độ đầy đủ của dữ liệu tương tác.",
            "Tập kiểm thử tích hợp và end-to-end chưa bao phủ toàn bộ luồng nghiệp vụ.",
            "Một số màn hình hoặc dữ liệu minh họa vẫn cần nối hoàn toàn với API sản xuất.",
            "Mô hình hiện tập trung vào Logistic Regression và Random Forest, chưa có quy trình tối ưu siêu tham số quy mô lớn.",
            "Cần bổ sung theo dõi drift và lịch tái huấn luyện khi triển khai dài hạn.",
        ],
    )
    add_heading(doc, "5.6. Đánh giá chung")
    add_text(doc, "Hệ thống đáp ứng mục tiêu học phần khi thể hiện việc sử dụng Python để tổ chức API, mô hình hóa nghiệp vụ, xử lý dữ liệu và xây dựng mô hình dự đoán. Sản phẩm có khả năng trình diễn rõ ràng và có nền tảng kỹ thuật để tiếp tục hoàn thiện với dữ liệu thực tế.")

    # Body page 16: conclusion and references.
    add_chapter(doc, "CHƯƠNG 6. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    add_heading(doc, "6.1. Kết luận")
    add_text(doc, "Nhóm đã xây dựng hệ thống phân tích dữ liệu khách hàng và dự đoán tiềm năng mua hàng gồm frontend, backend, cơ sở dữ liệu và các dịch vụ tích hợp. Đề tài kết hợp điểm tiềm năng dựa trên RFM và tương tác với mô hình dự đoán mua lại; kết quả được đưa vào dashboard và danh sách ưu tiên để hỗ trợ người quản trị.")
    add_text(doc, "Quá trình thực hiện giúp nhóm vận dụng Python vào một hệ thống có nhiều lớp, sử dụng lập trình hướng đối tượng, bất đồng bộ, kiểm tra dữ liệu và kiểm thử. Các thành phần được tách theo trách nhiệm, tạo điều kiện cho việc bảo trì và mở rộng sau môn học.")
    add_heading(doc, "6.2. Hướng phát triển")
    add_bullets(
        doc,
        [
            "Bổ sung dữ liệu thực tế và đánh giá mô hình theo nhiều giai đoạn thời gian.",
            "Tự động theo dõi chất lượng dữ liệu, drift và lịch tái huấn luyện.",
            "Mở rộng chiến dịch chăm sóc theo phân khúc và đo lường hiệu quả sau hành động.",
            "Hoàn thiện kiểm thử tích hợp, end-to-end và quy trình triển khai tự động.",
            "Bổ sung giải thích mô hình để người quản trị hiểu yếu tố ảnh hưởng đến dự đoán.",
        ],
    )
    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        (
            "[1] Nhóm 6. Mã nguồn Customer Analytics Platform, 2026. ",
            "GitHub repository",
            "https://github.com/phh235/customer-analytics-platform",
        ),
        (
            "[2] Python Software Foundation. Python 3.13 Documentation. ",
            "https://docs.python.org/3/",
            "https://docs.python.org/3/",
        ),
        (
            "[3] FastAPI. FastAPI Documentation. ",
            "https://fastapi.tiangolo.com/",
            "https://fastapi.tiangolo.com/",
        ),
        (
            "[4] SQLAlchemy. SQLAlchemy 2.0 Documentation. ",
            "https://docs.sqlalchemy.org/",
            "https://docs.sqlalchemy.org/",
        ),
        (
            "[5] PostgreSQL Global Development Group. PostgreSQL Documentation. ",
            "https://www.postgresql.org/docs/",
            "https://www.postgresql.org/docs/",
        ),
        (
            "[6] L. Breiman. Random Forests. Machine Learning, 45, 5-32, 2001. ",
            "https://doi.org/10.1023/A:1010933404324",
            "https://doi.org/10.1023/A:1010933404324",
        ),
        (
            "[7] T. Fawcett. An Introduction to ROC Analysis. Pattern Recognition Letters, 27(8), 861-874, 2006. ",
            "https://doi.org/10.1016/j.patrec.2005.10.010",
            "https://doi.org/10.1016/j.patrec.2005.10.010",
        ),
    ]
    for prefix, link_text, url in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        set_font(p.add_run(prefix), size=10.5)
        add_hyperlink(p, link_text, url)

    # Document settings and metadata.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.core_properties.title = "Hệ thống phân tích dữ liệu khách hàng và dự đoán tiềm năng mua hàng"
    doc.core_properties.subject = "Báo cáo đồ án môn Kỹ thuật lập trình Python"
    doc.core_properties.author = "Nhóm 6"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
