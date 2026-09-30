"""Finalize the project report from the latest complete Word revision."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(r"D:\Monorepo\customer-analytics-platform\output\documents")
SOURCE = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_final_v9.docx"
OUTPUT = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_hoan_chinh_v2.docx"

REFERENCES = [
    ("Python Software Foundation. Python 3.13 Documentation.", "https://docs.python.org/3/"),
    ("FastAPI. FastAPI Documentation.", "https://fastapi.tiangolo.com/"),
    ("Pydantic. Pydantic Documentation.", "https://docs.pydantic.dev/latest/"),
    ("SQLAlchemy. SQLAlchemy Documentation.", "https://docs.sqlalchemy.org/"),
    ("Alembic. Alembic Documentation.", "https://alembic.sqlalchemy.org/en/latest/"),
    ("PostgreSQL Global Development Group. PostgreSQL Documentation.", "https://www.postgresql.org/docs/"),
    ("Neon. Neon Documentation.", "https://neon.com/docs"),
    ("React. React Documentation.", "https://react.dev/learn"),
    ("Microsoft. TypeScript Documentation.", "https://www.typescriptlang.org/docs/"),
    ("Vite. Vite Guide.", "https://vite.dev/guide/"),
    ("Tailwind Labs. Tailwind CSS Documentation.", "https://tailwindcss.com/docs/installation/using-vite"),
    ("shadcn/ui. Component Documentation.", "https://ui.shadcn.com/docs"),
    ("TanStack. TanStack Query React Documentation.", "https://tanstack.com/query/latest/docs/framework/react/overview"),
    ("Cloudinary. Cloudinary Documentation.", "https://cloudinary.com/documentation"),
    ("Groq. Groq API Documentation.", "https://console.groq.com/docs/overview"),
    ("Vitest. Vitest Guide.", "https://vitest.dev/guide/"),
    ("L. Breiman. Random Forests. Machine Learning, 45, 5–32, 2001.", "https://doi.org/10.1023/A:1010933404324"),
    ("T. Fawcett. An Introduction to ROC Analysis. Pattern Recognition Letters, 27(8), 861–874, 2006.", "https://doi.org/10.1016/j.patrec.2005.10.010"),
    ("skills.sh. The Agent Skills Directory.", "https://www.skills.sh/"),
    ("OMP. Coding agent with IDE integration.", "https://omp.sh/"),
    ("Orca. Agent development environment.", "https://www.onorca.dev/"),
    ("OpenAI. Codex.", "https://openai.com/codex/"),
]


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for extra in paragraph.runs[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        paragraph.add_run(value)


def add_link(paragraph, url, label=None):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    properties.extend((color, underline, size))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label or url
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


document = Document(SOURCE)


def cite_paragraph(prefix, citation):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    assert len(matches) == 1, f"Expected one paragraph beginning {prefix!r}"
    matches[0].add_run(f" {citation}")


def cite_fragment(prefix, fragment, replacement):
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith(prefix)]
    assert len(matches) == 1, f"Expected one paragraph beginning {prefix!r}"
    for run in matches[0].runs:
        if fragment in run.text:
            run.text = run.text.replace(fragment, replacement, 1)
            return
    raise AssertionError(f"Missing fragment {fragment!r}")

# Reflect the actual frontend and hosted PostgreSQL stack in the summary table.
technology = next(table for table in document.tables if table.cell(0, 0).text == "Nhóm")
set_cell_text(technology.cell(1, 1), "React, TypeScript, Vite, Tailwind CSS, shadcn/ui")
set_cell_text(technology.cell(3, 1), "PostgreSQL trên Neon, SQLAlchemy, Alembic")

# Point readers from the technical narrative to the numbered source list.
cite_paragraph("Python 3.13 là ngôn ngữ chính", "[1]–[4], [6]")
cite_fragment("Mô hình sử dụng các đặc trưng như recency", "Random Forest", "Random Forest [17]")
cite_paragraph("Mô hình được huấn luyện từ cửa sổ đặc trưng", "Hệ thống còn lưu ROC-AUC như một chỉ số tham khảo khi đánh giá [18].")
cite_fragment("Frontend gọi các API có kiểu dữ liệu rõ ràng", "TanStack Query.", "TanStack Query [13].")
cite_fragment("API sản phẩm trả danh sách phân trang", "Cloudinary.", "Cloudinary [14].")
cite_fragment("Trợ lý 3CS AI sử dụng Groq Cloud", "Groq Cloud", "Groq Cloud [15]")
cite_fragment("Frontend sử dụng TypeScript để kiểm tra kiểu", "TypeScript để", "TypeScript [9] để")
cite_fragment("Frontend sử dụng TypeScript [9]", "Vitest cho", "Vitest [16] cho")

source_note = document.add_paragraph(
    "Nguồn tài liệu cho giao diện: React [8], TypeScript [9], Vite [10], "
    "Tailwind CSS [11] và shadcn/ui [12]."
)
source_note.paragraph_format.space_before = Pt(4)
source_note.paragraph_format.space_after = Pt(6)
source_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in source_note.runs:
    run.font.size = Pt(10)
technology._tbl.addnext(source_note._p)

# Keep the technical description of the system, but omit a repository URL from
# the bibliography because the source code is submitted separately as a ZIP.
reference_heading = next(
    paragraph for paragraph in document.paragraphs if paragraph.text == "TÀI LIỆU THAM KHẢO"
)
old_references = []
found_heading = False
for paragraph in document.paragraphs:
    if paragraph._p is reference_heading._p:
        found_heading = True
        continue
    if found_heading and paragraph.text.lstrip().startswith("["):
        old_references.append(paragraph)
for paragraph in old_references:
    paragraph._element.getparent().remove(paragraph._element)
assert len(old_references) == 7, f"Expected 7 old references, got {len(old_references)}"

for index, (description, url) in enumerate(REFERENCES, start=1):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.paragraph_format.left_indent = Pt(17)
    paragraph.paragraph_format.first_line_indent = Pt(-17)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    prefix = paragraph.add_run(f"[{index}] {description} ")
    prefix.font.size = Pt(10.5)
    display_url = {
        "https://tanstack.com/query/latest/docs/framework/react/overview": "https://tanstack.com/query",
        "https://tailwindcss.com/docs/installation/using-vite": "https://tailwindcss.com/docs",
    }.get(url, url)
    add_link(paragraph, url, display_url)

# Drop the now-unused relationship to the separately packaged source code.
for key, relationship in list(document.part.rels.items()):
    if "github.com/phh235/customer-analytics-platform" in relationship.target_ref:
        del document.part.rels[key]

document.save(OUTPUT)
print(f"Updated {OUTPUT} with {len(REFERENCES)} clickable references")
