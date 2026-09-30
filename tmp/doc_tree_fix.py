"""Use plain Vietnamese labels in the report's project structure tree."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(r"D:\Monorepo\customer-analytics-platform\output\documents")
SOURCE = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_hoan_chinh_v2.docx"
OUTPUT = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_hoan_chinh_v3.docx"

document = Document(SOURCE)
paragraphs = [
    p for p in document.paragraphs if p.text.startswith("customer-analytics-platform/")
]
assert len(paragraphs) == 1
paragraph = paragraphs[0]
paragraph.clear()
run = paragraph.add_run(
    "customer-analytics-platform/\n"
    "├── frontend/src/        # Giao diện và kết nối API\n"
    "├── backend/src/customer_analytics/\n"
    "│   ├── app/features/    # Các nhóm chức năng\n"
    "│   ├── app/shared/      # Thành phần dùng chung\n"
    "│   └── core/            # Cơ sở dữ liệu và hạ tầng\n"
    "├── backend/migrations/  # Cập nhật cấu trúc CSDL\n"
    "└── backend/tests/       # Kiểm thử backend"
)
run.font.name = "Consolas"
run.font.size = Pt(9.5)
document.save(OUTPUT)
print(OUTPUT)
