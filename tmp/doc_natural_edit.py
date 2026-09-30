"""Polish the report's Vietnamese while preserving technical meaning and layout."""

from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Monorepo\customer-analytics-platform\output\documents")
SOURCE = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_hoan_chinh_v3.docx"
OUTPUT = ROOT / "Do_an_Ky_thuat_lap_trinh_Python_Nhom_6_hoan_chinh_v4.docx"
document = Document(SOURCE)


def rewrite(index: int, expected_start: str, new_text: str) -> None:
    paragraph = document.paragraphs[index]
    assert paragraph.text.startswith(expected_start), (index, paragraph.text[:80])
    assert not paragraph._p.xpath(".//w:hyperlink"), index
    assert paragraph.runs, index
    paragraph.runs[0].text = new_text
    for extra in paragraph.runs[1:]:
        extra._element.getparent().remove(extra._element)


REWRITES = {
    80: ("Nhóm 6 xin chân thành", "Nhóm 6 cảm ơn ThS. Nghi Hoàng Khoa đã hướng dẫn và góp ý trong quá trình thực hiện đồ án môn Kỹ thuật lập trình Python. Những góp ý của thầy giúp nhóm điều chỉnh cách tổ chức chương trình và hoàn thiện báo cáo."),
    81: ("Nhóm cũng cảm ơn", "Nhóm cảm ơn Trường Đại học Công nghệ Thông tin đã tạo điều kiện học tập. Do thời gian thực hiện và dữ liệu thử nghiệm còn hạn chế, báo cáo có thể còn thiếu sót; nhóm mong nhận được góp ý để tiếp tục cải thiện."),
    84: ("Hoạt động bán hàng", "Trong quá trình bán hàng, thông tin về khách hàng, đơn hàng, sản phẩm và lượt xem thường nằm ở nhiều nơi. Khi xem từng nguồn riêng lẻ, người quản trị khó biết khách hàng nào cần chăm sóc trước và ai có khả năng mua lại. Vì vậy, nhóm chọn xây dựng hệ thống phân tích dữ liệu khách hàng và dự đoán tiềm năng mua hàng trong phạm vi đồ án môn học."),
    85: ("Sản phẩm gồm website", "Hệ thống gồm trang cửa hàng cho khách hàng, trang quản trị và API. Backend Python tiếp nhận dữ liệu, tính điểm tiềm năng và dự đoán khả năng mua lại. Giao diện cho phép xem sản phẩm, ghi nhận lượt xem và theo dõi kết quả phân tích."),
    86: ("Báo cáo tập trung", "Báo cáo trình bày cách nhóm tổ chức backend Python, nhập và xử lý dữ liệu, tính điểm, huấn luyện mô hình và kiểm thử. Kết quả dự đoán được dùng để tham khảo khi chọn khách hàng cần ưu tiên; độ tin cậy của kết quả phụ thuộc vào chất lượng và lượng dữ liệu."),
    90: ("Dữ liệu khách hàng chỉ", "Khi hồ sơ khách hàng, đơn hàng và lượt xem sản phẩm nằm ở nhiều nơi, người quản trị phải đối chiếu từng phần để biết ai có khả năng mua lại. Nhóm kết hợp các dữ liệu đó để tính điểm, theo dõi xu hướng và lập danh sách khách hàng cần quan tâm. Bài toán này cũng giúp nhóm vận dụng Python trong xây dựng API và xử lý dữ liệu."),
    98: ("Đề tài tập trung", "Đề tài tập trung vào khách hàng, sản phẩm, đơn hàng và lượt xem sản phẩm. Dữ liệu được lưu trong PostgreSQL; hệ thống nhận file Excel theo mẫu đã quy định. Mô hình dự đoán khả năng mua lại trong khoảng thời gian được chọn. Một số màn hình vẫn dùng dữ liệu mẫu, nên kết quả hiện phù hợp để trình diễn và thử nghiệm."),
    104: ("Phân tích khách hàng kết hợp", "Nhóm dựa vào ba chỉ số RFM: thời gian từ lần mua gần nhất, số lần mua và tổng giá trị mua hàng. Lượt tương tác với sản phẩm được bổ sung để nhận biết sự quan tâm của khách hàng ngay cả khi họ chưa đặt đơn."),
    106: ("Mỗi thành phần được", "Mỗi chỉ số được chấm tối đa 5 điểm. Trọng số hiện dùng là Recency 35%, Frequency 30%, Monetary 20% và Interaction 15%. Tổng điểm được quy về thang 100:"),
    108: ("Hệ thống phân loại", "Từ 80 điểm trở lên là mức Cao; từ 60 đến dưới 80 là Tiềm năng; dưới 60 là Bình thường. Nếu khách hàng chưa có đơn hàng hợp lệ trong thời gian phân tích, hệ thống hiển thị thiếu dữ liệu thay vì tự gán 0 điểm."),
    110: ("Mô hình sử dụng", "Mô hình dựa vào thời gian mua gần nhất, số lần mua, tổng chi tiêu, giá trị đơn trung bình, chu kỳ mua, mức tương tác, độ đa dạng sản phẩm và điểm đánh giá. Backend hỗ trợ Logistic Regression và Random Forest [17], đều được viết bằng Python. Nhóm dùng PR-AUC, Lift@Top10 và Precision@Top10 để đánh giá khả năng xếp hạng khách hàng."),
    112: ("Mô hình được huấn luyện", "Khi huấn luyện, hệ thống lấy dữ liệu hành vi trong một khoảng thời gian trước mốc dự đoán và xác định khách hàng có mua lại sau mốc đó hay không. Mô hình mới chỉ được sử dụng khi đạt các ngưỡng đánh giá và được quản trị viên triển khai. Hệ thống cũng lưu ROC-AUC để tham khảo khi so sánh các phiên bản [18]."),
    115: ("Python 3.13 là", "Backend dùng Python 3.13 [1] và FastAPI [2] để xây dựng API; Pydantic [3] kiểm tra dữ liệu đầu vào. SQLAlchemy [4] giúp truy cập PostgreSQL [6] theo cách bất đồng bộ. Nhóm dùng lớp và giao diện để tách phần xử lý nghiệp vụ khỏi phần truy cập dữ liệu."),
    118: ("Nguồn tài liệu cho giao diện", "Tài liệu tham khảo cho phần giao diện: React [8], TypeScript [9], Vite [10], Tailwind CSS [11] và shadcn/ui [12]."),
    121: ("Luồng phân tích bắt đầu", "Dữ liệu khách hàng, đơn hàng và lượt tương tác được nhập vào hệ thống trước. Backend tổng hợp các chỉ số và tính điểm tiềm năng. Khi cần, quản trị viên huấn luyện và triển khai phiên bản mô hình mới để dự đoán khả năng mua lại. Kết quả được hiển thị trên dashboard và danh sách ưu tiên; phiên bản cấu hình và mô hình được lưu để kiểm tra lại."),
    135: ("Hệ thống cần phản hồi", "Backend cần báo lỗi rõ khi dữ liệu không hợp lệ, bảo vệ phiên đăng nhập và kiểm tra quyền ở từng API. Mật khẩu, token và thông tin cấu hình nhạy cảm không được ghi vào log hoặc mã nguồn. Giao diện cần dùng được trên màn hình nhỏ, có chế độ sáng/tối và hiển thị trạng thái chờ khi tải dữ liệu."),
    139: ("Frontend gọi các API", "Frontend gọi API và dùng TanStack Query [13] để cập nhật dữ liệu trên giao diện. Backend chia thành lớp tiếp nhận yêu cầu, lớp xử lý nghiệp vụ và lớp truy cập dữ liệu PostgreSQL. Cloudinary lưu ảnh sản phẩm, còn Groq phục vụ trợ lý 3CS AI."),
    152: ("Access token được", "Frontend giữ access token trong bộ nhớ, còn backend lưu refresh token trong cookie HTTP-only. Nếu nhiều yêu cầu cùng gặp lỗi 401, frontend chỉ gửi một yêu cầu làm mới phiên rồi thử lại. Giao diện chặn các trang không đúng vai trò; backend kiểm tra quyền khi xử lý API."),
    161: ("Sau khi cấu hình", "Sau khi thiết lập kết nối PostgreSQL trong backend/.env, nhóm chạy pnpm setup để cài các thư viện và pnpm dev để chạy frontend cùng backend. Trong quá trình phát triển, nhóm dùng trang /docs của FastAPI để kiểm tra API."),
    164: ("Backend kiểm tra email", "Backend kiểm tra email, mật khẩu và trạng thái tài khoản khi đăng nhập; các API tiếp tục kiểm tra quyền trước khi xử lý. Access token có thời hạn ngắn. Người dùng có thể đăng xuất hoặc đặt lại mật khẩu để thu hồi phiên. Chức năng quên mật khẩu dùng mã OTP sáu chữ số, có thời hạn và chỉ dùng một lần."),
    166: ("API sản phẩm trả", "API sản phẩm hỗ trợ xem danh sách theo trang, xem chi tiết và gợi ý tối đa bốn sản phẩm cùng danh mục. Ảnh được kiểm tra loại tệp, dung lượng rồi tải lên Cloudinary [14]. Khi khách hàng mở trang chi tiết, hệ thống lưu một lượt xem (product_view) để phân tích hành vi."),
    172: ("Trợ lý 3CS AI sử dụng", "Trợ lý 3CS AI dùng Groq Cloud [15] để trả lời câu hỏi về dữ liệu theo quyền của người dùng. Người dùng có thể mở khung trò chuyện trên trang quản trị, kể cả khi dùng màn hình nhỏ."),
    176: ("Use case tính điểm", "Chức năng tính điểm lấy trọng số từ cấu hình và tổng hợp dữ liệu giao dịch, tương tác trong thời gian được chọn. Hệ thống chấm từng thành phần rồi lưu điểm cùng phiên bản cấu hình. Nếu không có dữ liệu tương tác, trọng số còn lại được điều chỉnh; nếu thiếu dữ liệu giao dịch cần thiết, hệ thống báo thiếu dữ liệu."),
    178: ("Request huấn luyện", "Khi huấn luyện, quản trị viên chọn loại mô hình, thời gian lấy dữ liệu và khoảng thời gian muốn dự đoán. Nếu không chọn ngày phân tích, backend tự lấy mốc phù hợp từ dữ liệu đơn hàng. Sau huấn luyện, hệ thống lưu tệp mô hình và các chỉ số đánh giá theo phiên bản. Chỉ mô hình đã được phê duyệt mới có thể triển khai; khi chưa có mô hình triển khai, API dự đoán trả lỗi."),
    182: ("Danh sách ưu tiên kết hợp", "Danh sách ưu tiên dùng cả điểm tiềm năng và xác suất mua lại. Khách hàng đạt từ 80 điểm hoặc có xác suất mua lại vượt ngưỡng cấu hình sẽ xuất hiện trong danh sách, xếp theo mức ưu tiên để quản trị viên theo dõi."),
    187: ("Dashboard cho phép", "Dashboard có bộ lọc theo thời gian, phân khúc, mức tiềm năng, nhóm sản phẩm và nhân viên phụ trách. Người quản trị xem được doanh thu, đơn hàng, số khách hàng, giá trị đơn trung bình, biểu đồ xu hướng và thông tin về chất lượng dữ liệu."),
    194: ("Trong quá trình thực hiện", "Nhóm dùng OMP, Orca và Codex Pro 20x để hỗ trợ đọc mã, sửa lỗi, xây dựng giao diện, viết kiểm thử và soạn báo cáo. Nhóm xem lại kết quả và chạy các bước kiểm tra phù hợp trước khi đưa thay đổi vào đồ án."),
    200: ("Các skill được cài đặt", "Nhóm tải các skill từ skills.sh để tham khảo cách làm với FastAPI, React, TypeScript và shadcn/ui. Bảng dưới liệt kê các skill đã dùng."),
    207: ("Đồ án hoàn thành", "Nhóm hoàn thành các chức năng chính cho trang khách hàng và trang quản trị. Khách hàng có thể xem sản phẩm và đăng nhập; quản trị viên quản lý dữ liệu, xem báo cáo, phân khúc và danh sách ưu tiên, đồng thời huấn luyện mô hình. Backend cung cấp API cho xác thực, quản lý dữ liệu và phân tích."),
    212: ("Frontend sử dụng", "Frontend được kiểm tra bằng TypeScript [9], ESLint, Prettier và Vitest [16]. Backend được kiểm tra bằng Ruff, mypy và pytest. Kết quả chạy lại trên mã nguồn hiện tại: frontend đạt 113/113 bài kiểm thử trên 32 tệp; backend đạt 71/71 bài kiểm thử. Các bài kiểm thử tập trung vào nghiệp vụ, dữ liệu đầu vào, xử lý lỗi và API."),
    215: ("•  Kiến trúc backend", "•  Backend tách phần nghiệp vụ khỏi phần truy cập cơ sở dữ liệu, giúp kiểm thử từng phần."),
    216: ("•  Luồng xác thực", "•  Access token được giữ trong bộ nhớ; refresh token được đặt trong cookie HTTP-only."),
    217: ("•  Điểm tiềm năng", "•  Kết quả chấm điểm lưu cả điểm thành phần, trọng số và phiên bản cấu hình để kiểm tra lại."),
    218: ("•  Vòng đời mô hình", "•  Mô hình chỉ phục vụ dự đoán sau khi được đánh giá, phê duyệt và triển khai."),
    219: ("•  Giao diện hỗ trợ", "•  Giao diện dùng được trên màn hình nhỏ, có chế độ tối và trạng thái chờ khi tải dữ liệu."),
    221: ("•  Chất lượng dự đoán", "•  Dữ liệu đơn hàng hoặc tương tác quá ít sẽ làm kết quả dự đoán kém ổn định."),
    222: ("•  Tập kiểm thử tích hợp", "•  Kiểm thử tích hợp và kiểm thử toàn hệ thống chưa bao phủ hết các chức năng."),
    223: ("•  Một số màn hình", "•  Một số màn hình còn dùng dữ liệu mẫu, chưa lấy đầy đủ dữ liệu từ API."),
    224: ("•  Mô hình hiện tập trung", "•  Phần dự đoán mới hỗ trợ Logistic Regression và Random Forest; nhóm chưa tự động thử nhiều bộ siêu tham số."),
    225: ("•  Cần bổ sung theo dõi drift", "•  Chưa có chức năng theo dõi khi dữ liệu thay đổi so với lúc huấn luyện và nhắc lịch huấn luyện lại."),
    227: ("Hệ thống đáp ứng mục tiêu", "Đồ án đáp ứng yêu cầu học phần ở các phần xây dựng API bằng Python, xử lý dữ liệu và dự đoán khả năng mua lại. Các chức năng chính chạy được trong buổi demo. Trước khi dùng với dữ liệu thực tế, nhóm cần hoàn thiện các màn hình còn dùng dữ liệu mẫu và đánh giá mô hình trên tập dữ liệu lớn hơn."),
    231: ("Nhóm đã xây dựng", "Nhóm đã xây dựng trang khách hàng, trang quản trị và backend để quản lý dữ liệu, tính điểm tiềm năng và dự đoán khả năng mua lại. Dashboard và danh sách ưu tiên giúp quản trị viên xem kết quả theo từng khách hàng."),
    232: ("Quá trình thực hiện", "Qua đồ án, nhóm thực hành Python, lập trình hướng đối tượng, xử lý bất đồng bộ và kiểm thử trên cùng một hệ thống. Việc tách giao diện, xử lý nghiệp vụ và truy cập dữ liệu giúp nhóm dễ xác định phần cần sửa và viết kiểm thử."),
    234: ("•  Bổ sung dữ liệu thực tế", "•  Thu thập thêm dữ liệu thực tế và đánh giá mô hình ở các mốc thời gian khác nhau."),
    235: ("•  Tự động theo dõi", "•  Theo dõi chất lượng và sự thay đổi của dữ liệu đầu vào để quyết định khi nào cần huấn luyện lại."),
    236: ("•  Mở rộng chiến dịch", "•  Bổ sung chiến dịch chăm sóc theo phân khúc và theo dõi kết quả sau khi liên hệ khách hàng."),
    237: ("•  Hoàn thiện kiểm thử", "•  Bổ sung kiểm thử tích hợp, kiểm thử toàn hệ thống và tự động hóa quá trình triển khai."),
    238: ("•  Bổ sung giải thích", "•  Cho người quản trị xem những yếu tố ảnh hưởng nhiều nhất đến dự đoán."),
}

for index, (expected, replacement) in REWRITES.items():
    rewrite(index, expected, replacement)


def replace_heading(old: str, new: str, expected_count: int = 2) -> None:
    changed = 0
    for paragraph in document.paragraphs:
        if old in paragraph.text:
            for node in paragraph._p.xpath(".//w:t"):
                if old in (node.text or ""):
                    node.text = node.text.replace(old, new)
                    changed += 1
                    break
    assert changed == expected_count, (old, changed)


replace_heading(
    "4.11. Công cụ AI và Agent Skills hỗ trợ phát triển",
    "4.11. Công cụ AI và Agent Skills sử dụng",
)
replace_heading(
    "4.11.1. Nền tảng AI và môi trường agent",
    "4.11.1. Công cụ AI sử dụng",
)


LINKED_BULLETS = {
    196: "Hỗ trợ đọc, chỉnh sửa mã và xem lại các thay đổi ngay trong môi trường lập trình. ",
    197: "Giúp nhóm quản lý môi trường làm việc, chạy lệnh, mở trang web và đối chiếu các thay đổi. ",
    198: "Hỗ trợ viết, sửa và kiểm tra chương trình, đồng thời góp ý cho giao diện và báo cáo. ",
}
for index, replacement in LINKED_BULLETS.items():
    paragraph = document.paragraphs[index]
    assert len(paragraph.runs) == 2 and paragraph._p.xpath(".//w:hyperlink")
    paragraph.runs[1].text = replacement


def replace_cell(table_index: int, row: int, column: int, new_text: str) -> None:
    paragraph = document.tables[table_index].cell(row, column).paragraphs[0]
    assert paragraph.runs
    paragraph.runs[0].text = new_text
    for extra in paragraph.runs[1:]:
        extra._element.getparent().remove(extra._element)


# Align the member roles with the cover slide and use plain language in tables.
replace_cell(1, 1, 2, "Phân tích nghiệp vụ và xử lý dữ liệu")
replace_cell(1, 2, 2, "Phát triển giao diện và tích hợp API")
replace_cell(1, 3, 2, "Thiết kế và phát triển Backend API")
replace_cell(2, 1, 2, "Giao diện khách hàng và trang quản trị")
replace_cell(2, 2, 2, "Đăng nhập, dữ liệu và phân tích")
replace_cell(2, 3, 2, "Lưu dữ liệu và quản lý thay đổi cấu trúc bảng")
replace_cell(2, 4, 2, "Lưu ảnh sản phẩm và hỗ trợ trợ lý AI")
replace_cell(2, 5, 2, "Cài thư viện, chạy và kiểm thử")
replace_cell(3, 1, 1, "Xem sản phẩm và tìm hiểu thông tin trước khi mua")
replace_cell(3, 2, 1, "Quản lý dữ liệu, xem báo cáo và triển khai mô hình")
replace_cell(5, 5, 1, "Đọc file Excel, kiểm tra cấu trúc và lưu dữ liệu")
replace_cell(8, 0, 0, "Chỉ số")
replace_cell(9, 1, 1, "Tham khảo cách tổ chức API, kiểm tra dữ liệu và xử lý bất đồng bộ.")
replace_cell(9, 2, 1, "Tham khảo mẫu tổ chức backend và xử lý lỗi.")
replace_cell(9, 3, 1, "Tham khảo cách xây dựng các thành phần giao diện với shadcn/ui.")
replace_cell(9, 4, 1, "Tham khảo cách định nghĩa kiểu dữ liệu cho API và giao diện.")
replace_cell(9, 5, 1, "Sắp xếp các thành phần React để dễ dùng lại.")
replace_cell(9, 6, 1, "Kiểm tra cách dùng hook và hiệu năng giao diện React.")
replace_cell(10, 5, 2, "Trả điểm tương ứng hoặc báo thiếu dữ liệu")
replace_cell(10, 6, 2, "Backend tự chọn ngày phân tích")

# The final row tested the presentation deck, not a function of the product.
test_table = document.tables[10]
assert test_table.rows[-1].cells[0].text == "Video và slide"
test_table._tbl.remove(test_table.rows[-1]._tr)

document.save(OUTPUT)
print(f"Edited {len(REWRITES)} paragraphs and the affected tables: {OUTPUT}")
