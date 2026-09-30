"""Contract parser and normalizers for the canonical dataset workbook."""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo

from customer_analytics.app.shared.errors import ErrorCode
from customer_analytics.app.shared.exceptions import AppException

BUSINESS_TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")
EXCEL_EPOCH = datetime(1899, 12, 30)
FORMULA_ERRORS = {"#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "#N/A"}

DATASET_SHEET_HEADERS: dict[str, tuple[str, ...]] = {
    "Customers": (
        "customer_id",
        "customer_unique_id",
        "customer_zip_code_prefix",
        "customer_city",
        "customer_state",
        "CustomerName",
        "Region",
        "ProvinceCity",
        "Email",
        "Phone",
        "RegisteredDate",
        "owner_id",
        "CustomerStatus",
        "CustomerSegment",
    ),
    "Products": (
        "product_id",
        "product_category_name",
        "product_name_lenght",
        "product_description_lenght",
        "product_photos_qty",
        "product_weight_g",
        "product_length_cm",
        "product_height_cm",
        "product_width_cm",
        "Description",
        "ProductCategory",
        "ListPrice",
        "Currency",
        "ProductStatus",
    ),
    "Orders": (
        "order_id",
        "customer_id",
        "order_status",
        "order_purchase_timestamp",
        "order_approved_at",
        "order_delivered_carrier_date",
        "order_delivered_customer_date",
        "order_estimated_delivery_date",
        "CustomerID",
        "SalesChannel",
        "PaymentMethod",
        "Region",
        "ProvinceCity",
        "owner_id",
        "OrderSubtotal",
        "FreightTotal",
        "DiscountTotal",
        "OrderTotal",
        "Currency",
        "IsValidForRFM",
    ),
    "OderItems": (
        "order_id",
        "order_item_id",
        "product_id",
        "seller_id",
        "shipping_limit_date",
        "price",
        "freight_value",
        "Quantity",
        "LineAmount",
        "OrderDetailID",
        "discount_value",
        "line_subtotal",
        "line_total",
    ),
    "Payments": (
        "order_id",
        "payment_sequential",
        "payment_type",
        "payment_installments",
        "payment_value",
        "PaymentMethodVN",
        "Currency",
    ),
    "Reviews": (
        "review_id",
        "order_id",
        "review_score",
        "review_comment_title",
        "review_comment_message",
        "review_creation_date",
        "review_answer_timestamp",
        "CustomerID",
        "Channel",
    ),
    "Sellers": (
        "seller_id",
        "seller_zip_code_prefix",
        "seller_city",
        "seller_state",
        "Region",
        "SellerName",
    ),
    "Geolocation": (
        "geolocation_zip_code_prefix",
        "geolocation_lat",
        "geolocation_lng",
        "geolocation_city",
        "geolocation_state",
        "Region",
    ),
    "DataDictionary": (
        "SheetName",
        "FieldName",
        "BusinessName",
        "Description",
        "DataType",
        "LengthFormat",
        "Required",
        "KeyType",
        "ReferenceField",
        "AllowedValues",
        "ValidationRule",
        "NullHandling",
        "TransformationRule",
        "ExampleValue",
        "SourceType",
    ),
    "Employees": (
        "employee_id",
        "employee_name",
        "department",
        "email",
        "phone",
        "region_scope",
        "status",
        "start_date",
    ),
    "Interactions": (
        "interaction_id",
        "customer_unique_id",
        "product_id",
        "interaction_type",
        "interaction_timestamp",
        "channel",
        "session_id",
        "campaign_id",
        "interaction_value",
        "interaction_result",
    ),
    "Campaigns": (
        "campaign_id",
        "campaign_name",
        "campaign_type",
        "start_date",
        "end_date",
        "channel",
        "target_segment",
        "status",
        "budget_vnd",
    ),
    "CustomerAnalytics": (
        "AnalysisDate",
        "CustomerID",
        "OwnerID",
        "RecencyDays",
        "Frequency",
        "Monetary",
        "AOV",
        "AvgPurchaseCycleDays",
        "AvgReviewScore",
        "InteractionScore",
        "RScore",
        "FScore",
        "MScore",
        "PotentialScore",
        "Segment",
        "ModelStatus",
    ),
    "ModelRuns": (
        "run_id",
        "run_timestamp",
        "analysis_date",
        "model_type",
        "model_version",
        "status",
        "record_count",
        "notes",
    ),
}

IMPORTABLE_SHEETS = tuple(
    name for name in DATASET_SHEET_HEADERS if name != "DataDictionary"
)


@dataclass(frozen=True)
class DatasetSheet:
    """Validated sheet data from the canonical workbook."""

    name: str
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class DatasetWorkbook:
    """Validated canonical workbook."""

    sheets: dict[str, DatasetSheet]

    @property
    def importable_rows(self) -> int:
        """Return the number of rows that participate in persistence."""
        return sum(len(self.sheets[name].rows) for name in IMPORTABLE_SHEETS)


def parse_dataset_workbook(file_content: bytes) -> DatasetWorkbook:
    """Parse and validate the exact workbook contract."""
    try:
        import openpyxl
    except ImportError:
        raise AppException(
            error_code=ErrorCode.INTERNAL_SERVER_ERROR,
            message="Thư viện openpyxl chưa được cài đặt.",
        ) from None

    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(file_content), read_only=True, data_only=True
        )
        actual_sheets = tuple(workbook.sheetnames)
        expected_sheets = tuple(DATASET_SHEET_HEADERS)
        missing = [name for name in expected_sheets if name not in actual_sheets]
        unexpected = [name for name in actual_sheets if name not in expected_sheets]
        if missing or unexpected:
            raise AppException(
                error_code=ErrorCode.VALIDATION_ERROR,
                message=(
                    "Workbook phải có đúng các sheet theo contract. "
                    f"Thiếu: {missing or 'không có'}. "
                    f"Không mong đợi: {unexpected or 'không có'}."
                ),
            )

        sheets: dict[str, DatasetSheet] = {}
        for sheet_name, expected_headers in DATASET_SHEET_HEADERS.items():
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = list(next(rows))
            except StopIteration:
                raise AppException(
                    error_code=ErrorCode.VALIDATION_ERROR,
                    message=f"Sheet '{sheet_name}' không có header.",
                ) from None
            while raw_headers and raw_headers[-1] in (None, ""):
                raw_headers.pop()
            headers = tuple(
                str(value).strip() if value is not None else "" for value in raw_headers
            )
            if headers != expected_headers:
                raise AppException(
                    error_code=ErrorCode.VALIDATION_ERROR,
                    message=(
                        f"Header sheet '{sheet_name}' không khớp contract. "
                        f"Expected: {list(expected_headers)}; actual: {list(headers)}."
                    ),
                )

            sheet_rows = []
            for raw_row in rows:
                values = tuple(raw_row[: len(headers)])
                if not any(value is not None and value != "" for value in values):
                    continue
                sheet_rows.append(
                    {
                        header: sanitize_formula_error(value)
                        for header, value in zip(headers, values, strict=False)
                    }
                )
            sheets[sheet_name] = DatasetSheet(
                name=sheet_name,
                headers=headers,
                rows=tuple(sheet_rows),
            )
        workbook.close()
        return DatasetWorkbook(sheets=sheets)
    except AppException:
        raise
    except Exception as exc:
        raise AppException(
            error_code=ErrorCode.VALIDATION_ERROR,
            message=f"Lỗi đọc workbook dataset: {exc!s}",
        ) from exc


def sanitize_formula_error(value: Any) -> Any:
    """Convert formula error tokens into missing values, never persisted text."""
    if isinstance(value, str) and value.strip().upper() in FORMULA_ERRORS:
        return None
    return value


def is_missing(value: Any) -> bool:
    """Return whether an Excel value is blank or a formula error."""
    return value is None or value == "" or sanitize_formula_error(value) is None


def clean_text(value: Any) -> str | None:
    """Normalize nullable text fields."""
    if is_missing(value):
        return None
    return str(value).strip() or None


def decimal_value(value: Any, default: Decimal = Decimal("0")) -> Decimal:
    """Parse a numeric workbook value without accepting formula errors."""
    if is_missing(value):
        return default
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"Giá trị tiền/số không hợp lệ: {value!r}") from None
    if not parsed.is_finite():
        raise ValueError(f"Giá trị tiền/số không hữu hạn: {value!r}")
    return parsed


def int_value(value: Any, default: int = 0) -> int:
    """Parse a workbook integer value."""
    if is_missing(value):
        return default
    try:
        return int(Decimal(str(value)))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError(f"Giá trị integer không hợp lệ: {value!r}") from None


def bool_value(value: Any, default: bool = False) -> bool:
    """Parse common Excel boolean representations."""
    if is_missing(value):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    normalized = str(value).strip().upper()
    if normalized in {"TRUE", "YES", "Y", "1", "VALID", "ACTIVE"}:
        return True
    if normalized in {"FALSE", "NO", "N", "0", "INVALID", "INACTIVE"}:
        return False
    raise ValueError(f"Giá trị boolean không hợp lệ: {value!r}")


def datetime_value(value: Any, *, date_only: bool = False) -> datetime | date | None:
    """Parse Excel serials and ISO/date strings in the business timezone."""
    if is_missing(value):
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not math.isfinite(float(value)):
            raise ValueError(f"Ngày Excel không hợp lệ: {value!r}")
        parsed = EXCEL_EPOCH + timedelta(days=float(value))
    else:
        raw = str(value).strip()
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        except ValueError:
            parsed = None
            for fmt in (
                "%Y-%m-%d",
                "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y",
                "%d-%m-%Y",
                "%m/%d/%Y",
            ):
                try:
                    parsed = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                try:
                    numeric = float(raw)
                except ValueError:
                    raise ValueError(f"Ngày không hợp lệ: {value!r}") from None
                parsed = EXCEL_EPOCH + timedelta(days=numeric)

    if parsed is None:
        raise ValueError(f"Ngày không hợp lệ: {value!r}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=BUSINESS_TIMEZONE)
    if date_only:
        return parsed.date()
    return parsed


def canonical_currency(value: Any) -> str:
    """Normalize workbook currency labels."""
    normalized = clean_text(value) or "VND"
    return "VND" if normalized.upper() in {"VNĐ", "VND"} else normalized.upper()


def canonical_order_status(value: Any) -> str:
    """Normalize dataset statuses to the backend order contract."""
    normalized = (clean_text(value) or "COMPLETED").upper()
    return {
        "DELIVERED": "DELIVERED",
        "RETURNED": "REFUNDED",
        "CANCELED": "CANCELLED",
        "CANCELLED": "CANCELLED",
        "PROCESSING": "PROCESSING",
    }.get(normalized, normalized)
